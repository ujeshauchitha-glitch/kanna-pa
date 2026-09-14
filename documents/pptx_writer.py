"""Render a `Document` to a .pptx slide deck via `python-pptx`.

Each top-level `Section` becomes one slide: its `heading` is the slide
title, its `bullets` (preferred) or `paragraphs` become the body text,
and a `table`, if present, is added as its own shape on the same slide.
Slides are inherently flat, so `Section.level` (meaningful for DOCX/PDF
subsections) is not used here — every section gets one slide regardless
of nesting depth.
"""
from __future__ import annotations

from pathlib import Path

from core.errors import DocumentGenerationUnavailable
from documents.model import Document


def render_pptx(document: Document, path: str | Path) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise DocumentGenerationUnavailable(
            "the 'python-pptx' package is not installed; run `pip install kanna[documents]`"
        ) from exc

    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = document.title
    subtitle_text = document.subtitle or (f"By {document.author}" if document.author else "")
    if subtitle_text and len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = subtitle_text

    body_layout = prs.slide_layouts[1]  # "Title and Content" — has title + body placeholders
    title_only_layout = prs.slide_layouts[5] if len(prs.slide_layouts) > 5 else body_layout
    # "Title Only" — has a title placeholder but no body one, so a table-only section (no
    # bullets/paragraphs) still gets its heading instead of landing on a titleless Blank layout.

    for section in document.sections:
        lines = section.bullets or section.paragraphs
        has_text_placeholder = bool(lines)
        slide = prs.slides.add_slide(body_layout if has_text_placeholder else title_only_layout)

        if slide.shapes.title is not None:
            slide.shapes.title.text = section.heading or ""

        table_top = Inches(1.8)
        if has_text_placeholder and len(slide.placeholders) > 1:
            text_frame = slide.placeholders[1].text_frame
            text_frame.clear()
            text_frame.text = lines[0]
            for line in lines[1:]:
                paragraph = text_frame.add_paragraph()
                paragraph.text = line
            table_top = Inches(1.8) + Inches(0.35) * min(len(lines), 6)

        if section.table is not None and section.table.headers:
            _add_table(slide, section.table, top=table_top)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


def _add_table(slide, table_data, *, top) -> None:
    from pptx.util import Inches

    n_cols = len(table_data.headers)
    n_rows = len(table_data.rows) + 1
    height = Inches(0.4) * n_rows
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, Inches(0.5), top, Inches(9), height)
    table = graphic_frame.table

    for col, header in enumerate(table_data.headers):
        table.cell(0, col).text = header
    for row_idx, row in enumerate(table_data.rows, start=1):
        for col_idx, value in enumerate(row):
            if col_idx < n_cols:
                table.cell(row_idx, col_idx).text = str(value)

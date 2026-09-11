"""Tests for the assignment_solve tool.

Uses FakeProvider to script LLM responses; no network required.
"""
from __future__ import annotations

import json

import pytest

from core.agent.loop import AgentLoop
from core.agent.state import AgentState
from core.bootstrap import build_registry
from core.llm.base import LLMResponse
from core.llm.fake import FakeProvider
from core.permissions.gate import PreApprovedGate
from core.planner.llm_planner import LLMPlanner
from core.tools.registry import ToolRegistry
from tools.authoring.assignment import (
    AssignmentResult, Answer, Question,
    extract_questions, answer_questions, solve_assignment,
)
from tools.authoring.tools import AssignmentSolveTool
from tools.authoring.llm_author import LLMAuthor


# --- Helpers ---


def _fake_extract_response(title="Assignment", questions=None):
    payload = {
        "title": title,
        "questions": questions or [
            {"id": "Q1", "text": "What is 2+2?"},
            {"id": "Q2", "text": "What is the boiling point of water?"},
        ],
    }
    return LLMResponse(content=json.dumps(payload))


def _fake_answer_response(answers=None):
    payload = {
        "answers": answers or [
            {"question_id": "Q1", "answer_text": "4",
             "source_references": ["Basic arithmetic"], "unresolved": []},
            {"question_id": "Q2", "answer_text": "100C",
             "source_references": ["Physics knowledge"], "unresolved": []},
        ],
    }
    return LLMResponse(content=json.dumps(payload))


def _make_tool(responses=None):
    """Create an AssignmentSolveTool backed by FakeProvider with scripted responses."""
    if responses is None:
        responses = [_fake_extract_response(), _fake_answer_response()]
    provider = FakeProvider(responses)
    return AssignmentSolveTool(provider=provider), provider


def _assignment_registry(responses=None):
    """Build a registry with assignment_solve backed by FakeProvider."""
    registry = build_registry()
    tool, _ = _make_tool(responses)
    registry._tools["assignment_solve"] = tool
    return registry


# --- Unit tests for solve_assignment ---


def test_extract_questions():
    provider = FakeProvider([_fake_extract_response()])
    title, questions = extract_questions(provider, "source text", "solve")
    assert title == "Assignment"
    assert len(questions) == 2
    assert questions[0].id == "Q1"
    assert questions[0].text == "What is 2+2?"
    assert questions[1].id == "Q2"
    assert "boiling" in questions[1].text


def test_answer_questions():
    provider = FakeProvider([LLMResponse(content=json.dumps({
        "answers": [
            {"question_id": "Q1", "answer_text": "4",
             "source_references": ["Basic arithmetic"], "unresolved": []},
        ],
    }))])
    questions = [Question(id="Q1", text="What is 2+2?")]
    answers = answer_questions(provider, "source text", questions)
    assert len(answers) == 1
    assert answers[0].question_id == "Q1"
    assert answers[0].answer_text == "4"


def test_solve_assignment_full():
    provider = FakeProvider([_fake_extract_response(), _fake_answer_response()])
    result = solve_assignment(provider, "source text", "solve assignment")
    assert isinstance(result, AssignmentResult)
    assert result.title == "Assignment"
    assert len(result.questions) == 2
    assert len(result.answers) == 2
    assert result.provenance[0].startswith("source:")
    assert any("questions extracted: 2" in p for p in result.provenance)
    assert any("questions answered: 2" in p for p in result.provenance)


def test_solve_assignment_empty_questions():
    provider = FakeProvider([LLMResponse(content=json.dumps({
        "title": "Assignment", "questions": [],
    }))])
    result = solve_assignment(provider, "source text", "solve")
    assert len(result.questions) == 0
    assert len(result.answers) == 0
    assert any("No questions" in u for u in result.unresolved)


# --- Tool tests ---


def test_assignment_solve_tool_success(ctx):
    tool, _ = _make_tool()
    result = tool.execute({
        "source_text": "2+2=4. Water boils at 100C.",
        "task": "Answer the questions.",
    }, ctx)
    assert result.success
    assert result.data["title"] == "Assignment"
    assert result.data["questions_extracted"] == 2
    assert result.data["questions_answered"] == 2
    assert len(result.data["sections"]) == 2
    assert len(result.data["answers"]) == 2
    assert isinstance(result.data["unresolved"], list)
    assert isinstance(result.data["provenance"], list)


def test_assignment_solve_empty_source(ctx):
    tool, _ = _make_tool()
    result = tool.execute({"source_text": "", "task": "solve"}, ctx)
    assert not result.success
    assert result.error.code == "empty_source"


def test_assignment_solve_empty_task(ctx):
    tool, _ = _make_tool()
    result = tool.execute({"source_text": "some text", "task": ""}, ctx)
    assert not result.success
    assert result.error.code == "empty_task"


def test_assignment_solve_malformed_llm_response(ctx):
    provider = FakeProvider([
        LLMResponse(content="not json"),
        LLMResponse(content="also not json"),
    ])
    tool = AssignmentSolveTool(provider=provider)
    result = tool.execute({"source_text": "source", "task": "solve"}, ctx)
    assert not result.success
    assert result.error.code == "assignment_failed"


def test_assignment_solve_with_unresolved(ctx):
    provider = FakeProvider([
        _fake_extract_response(questions=[{"id": "Q1", "text": "Complex question?"}]),
        _fake_answer_response(answers=[
            {"question_id": "Q1", "answer_text": "Partial answer.",
             "source_references": [], "unresolved": ["Need more data."]},
        ]),
    ])
    tool = AssignmentSolveTool(provider=provider)
    result = tool.execute({"source_text": "partial info", "task": "solve"}, ctx)
    assert result.success
    assert result.data["unresolved"] == ["Need more data."]
    assert result.data["answers"][0]["unresolved"] == ["Need more data."]


def test_assignment_solve_tool_registered_in_bootstrap(ctx):
    registry = build_registry()
    assert registry.has("assignment_solve")


# --- Integration: end-to-end read DOCX -> assignment solve -> generate DOCX ---


def test_end_to_end_docx_assignment_to_docx(ctx, tmp_path):
    pytest.importorskip("docx")
    from documents.reader import DocumentReadTool
    from documents import reader as document_reader
    from tools import filesystem, process
    from documents import tools as document_tools

    # Create input DOCX.
    import docx
    doc = docx.Document()
    doc.add_heading("Math Assignment", level=0)
    doc.add_heading("Q1", level=1)
    doc.add_paragraph("What is 2+2?")
    doc.add_heading("Q2", level=1)
    doc.add_paragraph("What is 3*4?")
    doc.save(str(tmp_path / "assignment.docx"))

    # Build registry with real document_read + fake assignment_solve.
    registry = ToolRegistry(gate=PreApprovedGate({"fs_write_file", "document_generate_docx"}))
    filesystem.register_all(registry)
    process.register_all(registry)
    document_tools.register_all(registry)
    document_reader.register_all(registry)

    extract_resp = LLMResponse(content=json.dumps({
        "title": "Math Assignment",
        "questions": [
            {"id": "Q1", "text": "What is 2+2?"},
            {"id": "Q2", "text": "What is 3*4?"},
        ],
    }))
    answer_resp = LLMResponse(content=json.dumps({
        "answers": [
            {"question_id": "Q1", "answer_text": "4", "source_references": [], "unresolved": []},
            {"question_id": "Q2", "answer_text": "12", "source_references": [], "unresolved": []},
        ],
    }))
    assignment_provider = FakeProvider([extract_resp, answer_resp])
    registry._tools["assignment_solve"] = AssignmentSolveTool(provider=assignment_provider)

    steps = [
        {"tool_name": "document_read", "args": {"path": "assignment.docx"},
         "expected": {"success": True}},
        {"tool_name": "assignment_solve", "args": {
            "source_text": {"$ref": "0.data.sections"},
            "task": "Answer all math questions.",
        }, "expected": {"success": True, "data_nonempty_key": "title"}},
        {"tool_name": "document_generate_docx", "args": {
            "path": "answers.docx",
            "title": {"$ref": "1.data.title"},
            "sections": {"$ref": "1.data.sections"},
        }, "expected": {"success": True, "min_files_created": 1, "files_exist": ["answers.docx"]}},
    ]
    planner_provider = FakeProvider([LLMResponse(content=json.dumps({"steps": steps}))])
    result = AgentLoop(registry, LLMPlanner(planner_provider), ctx).run("Solve assignment")
    assert result.state == AgentState.COMPLETE
    assert (tmp_path / "answers.docx").exists()
    # Verify the output DOCX has content.
    out_doc = docx.Document(str(tmp_path / "answers.docx"))
    text = "\n".join(p.text for p in out_doc.paragraphs)
    assert "4" in text
    assert "12" in text

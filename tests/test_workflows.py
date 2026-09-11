"""Real tools and SQLite, scripted reasoning only; no network or model claims."""
import json
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

from core.agent.loop import AgentLoop
from core.agent.state import AgentState
from core.bootstrap import build_registry
from core.llm.base import LLMResponse
from core.llm.fake import FakeProvider
from core.planner.llm_planner import LLMPlanner
from core.planner.plan import Plan, PlanStep
from core.tools.result import ToolResult
from tools.process.run_process import ProcessTool


def run_plan(ctx, steps, registry=None, corrections=2):
    provider = FakeProvider([LLMResponse(content=json.dumps({"steps": steps}))])
    return AgentLoop(registry or build_registry(), LLMPlanner(provider), ctx,
                     max_corrections=corrections).run("Execute this workflow")


def step(tool, args, **expected):
    return {"tool_name": tool, "args": args, "expected": expected or {"success": True}}


def ref(path):
    return {"$ref": path}


def test_natural_language_create_file(ctx, tmp_path):
    provider = FakeProvider([LLMResponse(content=json.dumps({"steps": [
        step("fs_write_file", {"path": "hello.txt", "content": "hello"},
             min_files_created=1, files_exist=["hello.txt"]),
    ]}))])
    result = AgentLoop(build_registry(), LLMPlanner(provider), ctx).run(
        "Create a text file containing hello.")
    assert result.state == AgentState.COMPLETE
    assert (tmp_path / "hello.txt").read_text() == "hello"
    assert str(tmp_path / "hello.txt") in result.message


def test_read_write_verify_and_persist_bindings(ctx, tmp_path):
    (tmp_path / "source.txt").write_text("Actual source, not invented")
    result = run_plan(ctx, [
        step("fs_read_file", {"path": "source.txt"}),
        step("fs_write_file", {"path": "copy.txt", "content": ref("0.data.content")}),
        step("fs_file_info", {"path": ref("1.files_created.0")}, data_equals={"exists": True}),
    ])
    assert result.state == AgentState.COMPLETE
    assert (tmp_path / "copy.txt").read_text() == "Actual source, not invented"
    rows = ctx.db.query("SELECT * FROM plan_steps ORDER BY step_index")
    persisted = json.loads(rows[1]["result"])
    assert persisted["resolved_args"]["content"] == "Actual source, not invented"
    assert persisted["attempts"][0]["verification_errors"] == []
    assert json.loads(rows[1]["args"])["content"] == ref("0.data.content")


def test_read_file_to_real_pdf(ctx, tmp_path):
    pytest.importorskip("reportlab")
    (tmp_path / "source.txt").write_text("Observed report source.")
    result = run_plan(ctx, [
        step("fs_read_file", {"path": "source.txt"}, data_equals={"truncated": False}),
        step("document_generate_pdf", {"path": "report.pdf", "title": "Source report",
             "sections": [{"heading": "Source", "paragraphs": [ref("0.data.content")]}]},
             files_exist=["report.pdf"], min_files_created=1),
        step("fs_file_info", {"path": ref("1.data.path")}, data_equals={"is_file": True}),
    ])
    assert result.state == AgentState.COMPLETE
    assert (tmp_path / "report.pdf").read_bytes().startswith(b"%PDF-")
    assert result.outcomes[1].step.args["sections"][0]["paragraphs"] == ["Observed report source."]


@pytest.mark.parametrize("reference", ["1.data.path", "2.data.path", "-1.data.path", "0", "0.__class__"])
def test_invalid_references_block_before_any_execution(ctx, reference):
    result = run_plan(ctx, [step("fs_write_file", {"path": "first.txt", "content": "first"}),
                            step("fs_read_file", {"path": ref(reference)})])
    assert result.state == AgentState.BLOCKED
    assert ctx.db.query("SELECT * FROM execution_log") == []


@pytest.mark.parametrize("reference", ["0.data.missing", "0.files_created.99"])
def test_missing_result_path_fails_without_downstream_invocation(ctx, reference):
    result = run_plan(ctx, [step("fs_file_info", {"path": "absent"}),
                            step("fs_read_file", {"path": ref(reference)})])
    assert result.state == AgentState.FAILED
    assert result.outcomes[-1].result.error.code == "binding_failed"
    assert result.outcomes[-1].attempts == 0
    assert len(ctx.db.query("SELECT * FROM execution_log")) == 1


def test_bound_values_still_validate_before_execution(ctx):
    result = run_plan(ctx, [step("fs_file_info", {"path": "absent"}),
                            step("fs_write_file", {"path": "bad", "content": ref("0.data.exists")})])
    assert result.state == AgentState.FAILED
    assert result.outcomes[-1].result.error.code == "invalid_input"
    assert result.outcomes[-1].attempts == 1


def test_bound_paths_cannot_escape_sandbox(ctx, tmp_path):
    (tmp_path / "source.txt").write_text(str(tmp_path.parent / "outside.txt"))
    result = run_plan(ctx, [step("fs_read_file", {"path": "source.txt"}),
                            step("fs_write_file", {"path": ref("0.data.content"), "content": "bad"})])
    assert result.state == AgentState.FAILED
    assert result.outcomes[-1].result.error.code == "sandbox_violation"
    assert result.outcomes[-1].attempts == 1


def test_approval_denial_never_retries_or_revises(ctx, tmp_path):
    (tmp_path / "old.txt").write_text("keep")
    result = run_plan(ctx, [step("fs_write_file", {"path": "old.txt", "content": "changed",
                                                       "overwrite": True})])
    assert result.state == AgentState.FAILED
    assert result.outcomes[0].attempts == 1
    assert (tmp_path / "old.txt").read_text() == "keep"


def test_successful_call_with_failed_postcondition_is_not_complete_or_replayed(ctx):
    result = run_plan(ctx, [step("fs_file_info", {"path": "missing"}, data_equals={"exists": True}),
                            step("fs_write_file", {"path": "never.txt", "content": "never"})])
    assert result.state == AgentState.FAILED
    assert result.outcomes[0].result.error.code == "verification_failed"
    assert result.outcomes[0].attempts == 1
    assert len(result.outcomes) == 1
    assert ctx.db.query_one("SELECT status FROM plans")["status"] == "failed"
    assert ctx.db.query_one("SELECT status FROM plan_steps WHERE step_index=0")["status"] == "failed"


def test_false_artifact_claim_is_detected(ctx, monkeypatch):
    registry = build_registry()
    monkeypatch.setattr(registry.get("fs_write_file"), "execute",
                        lambda args, ctx: ToolResult.ok({"path": "ghost"}, files_created=["ghost"]))
    result = run_plan(ctx, [step("fs_write_file", {"path": "ghost", "content": "hello"})], registry)
    assert result.state == AgentState.FAILED
    assert "output does not exist" in result.message


@pytest.mark.parametrize("expected", [{"success": False}, {"made_up": True}, {"files_exist": "x"},
                                      {"min_files_created": -1}])
def test_invalid_postconditions_block(ctx, expected):
    result = run_plan(ctx, [step("fs_file_info", {"path": "."}, **expected)])
    assert result.state == AgentState.BLOCKED


def test_process_nonzero_is_failure_with_real_diagnostics(ctx):
    registry = build_registry()
    # Use the actual current runtime; production allowlist remains unchanged.
    registry._tools["process_run"] = ProcessTool(frozenset({Path(sys.executable).name}))
    result = run_plan(ctx, [step("process_run", {"command": [sys.executable, "-c",
                                "import sys; print('diagnostic'); sys.exit(3)"]})], registry)
    assert result.state == AgentState.FAILED
    assert result.outcomes[0].result.data["stdout"].strip() == "diagnostic"
    assert result.outcomes[0].attempts == 1


def test_finance_summary_returns_actual_database_answer(ctx):
    result = run_plan(ctx, [step("finance_add_transaction", {"text": "I spent INR 340 on lunch"}),
                            step("finance_query", {"text": "How much did I spend this month?"})])
    assert result.state == AgentState.COMPLETE
    assert result.outcomes[1].result.data["totals"] == [{"currency": "INR", "amount_minor": 34000}]
    assert result.outcomes[1].result.data["message"] in result.message


def test_structured_extraction_to_document_keeps_tables(ctx, tmp_path):
    pytest.importorskip("docx")
    from docx import Document
    from vision.document.base import DocumentStructure, StructureSection, StructureTable
    from vision.document.fake import FakeDocumentProvider
    from vision.tools import VisionExtractStructureTool
    provider = FakeDocumentProvider(structure_extractions=[DocumentStructure(
        title="Assignment", sections=[StructureSection(heading="Observations", paragraphs=["Measured"],
        table=StructureTable(headers=["Trial", "Value"], rows=[["1", "12"]]))])])
    registry = build_registry()
    registry._tools["vision_extract_structure"] = VisionExtractStructureTool(provider)
    (tmp_path / "input.pdf").write_bytes(b"scripted vision input")
    result = run_plan(ctx, [step("vision_extract_structure", {"path": "input.pdf"}),
                            step("document_generate_docx", {"path": "out.docx",
                                 "title": ref("0.data.title"), "sections": ref("0.data.sections")})], registry)
    assert result.state == AgentState.COMPLETE
    document = Document(tmp_path / "out.docx")
    assert document.tables[0].cell(1, 1).text == "12"
    assert "Observations" in [p.text for p in document.paragraphs]


def test_recovery_passes_corrected_result_and_keeps_attempt_evidence(ctx, tmp_path):
    (tmp_path / "actual.txt").write_text("recovered content")
    provider = FakeProvider([
        LLMResponse(content=json.dumps({"steps": [
            step("fs_read_file", {"path": "wrong.txt"}),
            step("fs_write_file", {"path": "copy.txt", "content": ref("0.data.content")}),
        ]})),
        LLMResponse(content=json.dumps({"args": {"path": "actual.txt"}})),
    ])
    result = AgentLoop(build_registry(), LLMPlanner(provider), ctx).run("copy actual.txt")
    assert result.state == AgentState.COMPLETE
    assert (tmp_path / "copy.txt").read_text() == "recovered content"
    history = result.outcomes[0].history
    assert len(history) == 2
    assert history[0]["result"]["error"]["code"] == "not_found"
    assert history[1]["args"] == {"path": "actual.txt"}
    row = ctx.db.query_one("SELECT result FROM plan_steps WHERE step_index=0")
    assert json.loads(row["result"])["attempts"] == history


def test_bound_objects_are_copies_not_mutable_prior_evidence():
    from core.planner.bindings import resolve
    results = [{"data": {"sections": [{"paragraphs": ["original"]}]}}]
    resolved = resolve({"sections": ref("0.data.sections")}, results)
    resolved["sections"][0]["paragraphs"].append("mutated")
    assert results[0]["data"]["sections"][0]["paragraphs"] == ["original"]


@pytest.mark.parametrize("payload", [{"steps": [None]}, {"steps": "bad"},
                                     {"steps": [{"tool_name": []}]},
                                     {"steps": [{"tool_name": "fs_read_file", "args": []}]}])
def test_malformed_model_plans_block_without_crashing(ctx, payload):
    provider = FakeProvider([LLMResponse(content=json.dumps(payload))])
    result = AgentLoop(build_registry(), LLMPlanner(provider), ctx).run("read")
    assert result.state == AgentState.BLOCKED


def test_expected_files_must_be_in_sandbox(ctx, tmp_path):
    result = run_plan(ctx, [step("fs_file_info", {"path": "."},
                                files_exist=[str(tmp_path.parent)])])
    assert result.state == AgentState.FAILED
    assert "outside the sandbox" in result.message


@pytest.mark.parametrize("task_request, expected", [("read missing.txt", "failed"),
                                              ("compose a symphony", "failed"),
                                              ("list files in .", "succeeded")])
def test_scheduled_agent_outcome_is_honest(ctx, task_request, expected):
    from automation.scheduler.scheduler import Scheduler
    from automation.scheduler.store import SchedulerStore
    from core.planner.rule_based import RuleBasedPlanner
    from interfaces.cli.commands.scheduler import _on_due
    store = SchedulerStore(ctx.db)
    store.create("workflow", "once", {"run_at": "2000-01-01T00:00:00", "request": task_request})
    kanna = SimpleNamespace(agent_loop=lambda: AgentLoop(build_registry(), RuleBasedPlanner(), ctx))
    outcomes = Scheduler(store, on_due=lambda s: _on_due(s, kanna)).tick()
    assert outcomes[0]["status"] == expected
    assert ctx.db.query_one("SELECT status FROM jobs")["status"] == expected



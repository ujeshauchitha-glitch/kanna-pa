"""Tests for the document_read tool (DOCX structure extraction).

Uses real python-docx for DOCX creation; no network required.
"""
from __future__ import annotations

import pytest

from documents.reader import DocumentReadTool, _read_docx


def _create_docx(path, title="Test Document", sections=None):
    """Helper to create a DOCX file with structured content."""
    import docx
    doc = docx.Document()
    doc.add_heading(title, level=0)
    for section in sections or []:
        heading = section.get("heading")
        level = section.get("level", 1)
        if heading:
            doc.add_heading(heading, level=level)
        for para in section.get("paragraphs", []):
            doc.add_paragraph(para)
        for bullet in section.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")
        table_data = section.get("table")
        if table_data:
            headers = table_data["headers"]
            rows = table_data.get("rows", [])
            table = doc.add_table(rows=1 + len(rows), cols=len(headers))
            table.style = "Table Grid"
            for i, h in enumerate(headers):
                table.rows[0].cells[i].text = h
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    table.rows[r + 1].cells[c].text = str(val)
    doc.save(str(path))


# --- Unit tests ---


def test_read_docx_basic(tmp_path):
    path = tmp_path / "test.docx"
    _create_docx(path, sections=[
        {"heading": "Introduction", "level": 1, "paragraphs": ["Hello world."]},
        {"heading": "Details", "level": 2, "paragraphs": ["More info.", "Even more."]},
    ])
    result = _read_docx(path)
    assert result["title"] == "Test Document"
    assert len(result["sections"]) == 2
    assert result["sections"][0]["heading"] == "Introduction"
    assert result["sections"][0]["paragraphs"] == ["Hello world."]
    assert result["sections"][1]["heading"] == "Details"
    assert result["sections"][1]["level"] == 2


def test_read_docx_with_bullets(tmp_path):
    path = tmp_path / "bullets.docx"
    _create_docx(path, sections=[
        {"heading": "Items", "paragraphs": [], "bullets": ["A", "B", "C"]},
    ])
    result = _read_docx(path)
    assert result["sections"][0]["bullets"] == ["A", "B", "C"]


def test_read_docx_with_table(tmp_path):
    path = tmp_path / "table.docx"
    _create_docx(path, sections=[
        {"heading": "Data", "paragraphs": [],
         "table": {"headers": ["Name", "Value"], "rows": [["X", "10"], ["Y", "20"]]}},
    ])
    result = _read_docx(path)
    assert len(result["sections"]) == 1
    assert result["sections"][0]["table"]["headers"] == ["Name", "Value"]
    assert result["sections"][0]["table"]["rows"] == [["X", "10"], ["Y", "20"]]


def test_read_docx_empty(tmp_path):
    path = tmp_path / "empty.docx"
    import docx
    doc = docx.Document()
    doc.save(str(path))
    result = _read_docx(path)
    assert result["title"] == ""
    assert result["sections"] == []


# --- Tool tests ---


def test_document_read_tool_success(ctx, tmp_path):
    path = tmp_path / "assignment.docx"
    _create_docx(path, sections=[
        {"heading": "Q1", "paragraphs": ["What is 2+2?"]},
    ])
    tool = DocumentReadTool()
    result = tool.execute({"path": "assignment.docx"}, ctx)
    assert result.success
    assert result.data["title"] == "Test Document"
    assert len(result.data["sections"]) == 1
    assert result.data["sections"][0]["heading"] == "Q1"


def test_document_read_tool_not_found(ctx):
    tool = DocumentReadTool()
    result = tool.execute({"path": "nonexistent.docx"}, ctx)
    assert not result.success
    assert result.error.code == "not_found"


def test_document_read_tool_not_docx(ctx, tmp_path):
    (tmp_path / "file.txt").write_text("hello")
    tool = DocumentReadTool()
    result = tool.execute({"path": "file.txt"}, ctx)
    assert not result.success
    assert result.error.code == "unsupported_format"


def test_document_read_tool_sandbox_escape(ctx):
    tool = DocumentReadTool()
    result = tool.execute({"path": "../outside/test.docx"}, ctx)
    assert not result.success
    assert result.error.code == "sandbox_violation"


def test_document_read_tool_registered_in_bootstrap(ctx):
    from core.bootstrap import build_registry
    registry = build_registry()
    assert registry.has("document_read")


# --- End-to-end: read DOCX -> author content -> generate PDF ---


def test_end_to_end_read_docx_author_generate_pdf(ctx, tmp_path):
    pytest.importorskip("reportlab")
    from core.agent.loop import AgentLoop
    from core.agent.state import AgentState
    from core.llm.base import LLMResponse
    from core.llm.fake import FakeProvider
    from core.permissions.gate import PreApprovedGate
    from core.planner.llm_planner import LLMPlanner
    from core.tools.registry import ToolRegistry
    from tools import filesystem, process
    from documents import tools as document_tools

    path = tmp_path / "input.docx"
    _create_docx(path, title="Lab Notes", sections=[
        {"heading": "Observations", "paragraphs": ["Temperature was 25C."]},
    ])

    # Build registry with authoring tool injected.
    from tools.authoring.tools import AuthorContentTool
    from tools.authoring.llm_author import LLMAuthor
    from documents import reader as document_reader
    registry = ToolRegistry(gate=PreApprovedGate({"fs_write_file", "document_generate_pdf"}))
    filesystem.register_all(registry)
    process.register_all(registry)
    document_tools.register_all(registry)
    document_reader.register_all(registry)
    author_provider = FakeProvider([LLMResponse(content=__import__("json").dumps({
        "title": "Lab Report",
        "sections": [{"heading": "Summary", "level": 1,
                      "paragraphs": ["Temperature was 25C."], "bullets": []}],
        "unresolved": [],
    }))])
    registry._tools["author_content"] = AuthorContentTool(author=LLMAuthor(author_provider))

    steps = [
        {"tool_name": "document_read", "args": {"path": "input.docx"},
         "expected": {"success": True}},
        {"tool_name": "author_content", "args": {
            "source_text": {"$ref": "0.data.sections"},
            "task": "Write a lab report summary.",
        }, "expected": {"success": True, "data_nonempty_key": "title"}},
        {"tool_name": "document_generate_pdf", "args": {
            "path": "report.pdf",
            "title": {"$ref": "1.data.title"},
            "sections": {"$ref": "1.data.sections"},
        }, "expected": {"success": True, "min_files_created": 1, "files_exist": ["report.pdf"]}},
    ]
    provider = FakeProvider([LLMResponse(content=__import__("json").dumps({"steps": steps}))])
    result = AgentLoop(registry, LLMPlanner(provider), ctx).run("Read and report")
    assert result.state == AgentState.COMPLETE
    assert (tmp_path / "report.pdf").read_bytes().startswith(b"%PDF-")

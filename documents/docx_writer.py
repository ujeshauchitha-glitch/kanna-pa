"""Render a `Document` to a .docx file via `python-docx`.

Lazy-imports `python-docx` (extra `[documents]`) so nothing outside this
module needs it installed — mirrors the lazy-import discipline in
`core/llm/anthropic_provider.py` and `vision/_common.py`.
"""
from __future__ import annotations

from pathlib import Path

from core.errors import DocumentGenerationUnavailable
from documents.model import Document


def render_docx(document: Document, path: str | Path) -> None:
    try:
        import docx
    except ImportError as exc:
        raise DocumentGenerationUnavailable(
            "the 'python-docx' package is not installed; run `pip install kanna[documents]`"
        ) from exc

    out = docx.Document()

    out.add_heading(document.title, level=0)
    if document.subtitle:
        try:
            out.add_paragraph(document.subtitle, style="Subtitle")
        except KeyError:  # the 'Subtitle' style may be absent from a stripped-down template
            out.add_paragraph(document.subtitle)
    if document.author:
        out.add_paragraph(f"By {document.author}")

    for section in document.sections:
        if section.heading:
            out.add_heading(section.heading, level=max(1, min(section.level, 9)))
        for paragraph in section.paragraphs:
            out.add_paragraph(paragraph)
        for bullet in section.bullets:
            out.add_paragraph(bullet, style="List Bullet")
        if section.table is not None and section.table.headers:
            table = out.add_table(rows=1, cols=len(section.table.headers))
            table.style = "Table Grid"
            for cell, header in zip(table.rows[0].cells, section.table.headers):
                cell.text = header
            for row in section.table.rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(value)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    out.save(str(path))
